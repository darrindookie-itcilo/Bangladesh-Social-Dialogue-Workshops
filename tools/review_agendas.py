"""Correct agenda labels and explicit placeholders, keeping the supplied timetable."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
import hashlib, json

ROOT = Path(__file__).resolve().parents[1]

def cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(value)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)

def main():
    rows = json.loads((ROOT/'verification/materials.json').read_text())
    for record in (r for r in rows if r['kind'] == 'agenda'):
        file = ROOT/'dist'/record['path']
        record.setdefault('sourceSha256', record['sha256'])
        document = Document(file)
        revision = '3.1' if file.name.startswith('A3519988') else '3'
        heading = document.paragraphs[0]
        heading.runs[0].text = f'DRAFT PROGRAMME  ·  REVISION {revision}  ·  23 October 2026'
        for run in heading.runs[1:]:
            run.text = ''
        metadata = document.tables[0]
        if not any(row.cells[0].text == 'Preparation phase' for row in metadata.rows):
            row = deepcopy(metadata.rows[1]._tr)
            metadata.rows[1]._tr.addnext(row)
            added = metadata.rows[2]
            cell_text(added.cells[0], 'Preparation phase')
            cell_text(added.cells[1], 'Pre-course learning and reflection; dates and arrangements to be confirmed. Session 1 includes a review of this preparation.')
        for row in document.tables[1].rows[1:]:
            if len(set(id(cell._tc) for cell in row.cells)) > 1 and not row.cells[2].text.strip():
                cell_text(row.cells[2], 'To be confirmed')
        document.save(file)
        record['sha256'] = hashlib.sha256(file.read_bytes()).hexdigest()
        record['edits'] = 'Revision label and date corrected; preparation phase restored as to be confirmed; missing session leads explicitly marked to be confirmed. Supplied session times and content retained.'
    for record in (r for r in rows if r['kind'] == 'pack'):
        file = ROOT/'dist'/record['path']
        with ZipFile(file, 'w', ZIP_DEFLATED) as archive:
            for name in record['entries']:
                archive.write(file.parent/name, name)
        record['sha256'] = hashlib.sha256(file.read_bytes()).hexdigest()
    (ROOT/'verification/materials.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2)+'\n')

if __name__ == '__main__':
    main()
