"""Verify original-source fidelity and decrypt every protected teaching file."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
from cryptography.exceptions import InvalidTag
import base64,getpass,hashlib,json,re

ROOT=Path(__file__).resolve().parents[1]
DIST=ROOT/'dist'
def main():
    meta=json.loads((DIST/'protected/kdf.json').read_text())
    password=getpass.getpass('Facilitator password (hidden): ')
    username=input('Facilitator username: ').strip().lower()
    material=(username+'\0'+password).encode();del password
    key=PBKDF2HMAC(algorithm=hashes.SHA256(),length=32,salt=base64.b64decode(meta['salt']),iterations=meta['iterations']).derive(material)
    aes=AESGCM(key);cache={}
    def plain(value):
        assert value[:4]==b'BD26'
        return aes.decrypt(value[4:16],value[16:],None)
    def reconstruct(value):
        out=plain(value)
        if out.startswith(b'BDREC1\n'):
            r=json.loads(out[7:]);chunks=[]
            for name,offset,length in r['parts']:
                if name not in cache:cache[name]=plain((DIST/'protected/parts'/name).read_bytes())
                chunks.append(cache[name][offset:offset+length])
            out=b''.join(chunks);assert len(out)==r['size']
        return out
    manifest=json.loads(reconstruct((DIST/'protected/manifest.enc').read_bytes()))
    recipes=json.loads(plain((DIST/'protected/recipes.enc').read_bytes()))
    rows=json.loads((ROOT/'verification/materials.json').read_text())
    archives={}
    for code in ['A3519987','A3519988']:
        archives[code]=ZipFile(next((ROOT.parent/'upload').glob(code+'_Complete_Materials(2).zip')))
    for row in rows:
        path=row['path']
        if path.startswith('protected/'):
            value=reconstruct(base64.b64decode(recipes[path]))
            source=archives[path.split('/')[1][:8]].read(row['source'])
            assert value==source,path
            with ZipFile(BytesIO(value)) as z:assert z.testzip() is None
        else:
            value=(DIST/path).read_bytes()
            if row['kind']=='handout':assert value==archives[Path(path).name[:8]].read(row['source'])
            if row['kind']=='agenda':
                source=(ROOT.parent/'upload'/row['source']).read_bytes()
                if row.get('edits'):
                    assert hashlib.sha256(source).hexdigest()==row['sourceSha256']
                    from docx import Document
                    doc=Document(BytesIO(value));assert 'REVISION 2' not in doc.paragraphs[0].text
                    assert any(r.cells[0].text=='Preparation phase' for r in doc.tables[0].rows)
                else:assert value==source
            if row['kind']=='pack':
                with ZipFile(BytesIO(value)) as z:
                    assert sorted(z.namelist())==row['entries']
                    assert not any(re.search(r'_(S\d+|R[12])_',n) for n in z.namelist())
                    assert not any(n.lower().endswith('.pdf') for n in z.namelist())
                    for name in z.namelist():assert z.read(name)==(DIST/Path(path).parent/name).read_bytes()
        assert hashlib.sha256(value).hexdigest()==row['sha256'],path
    assert len(manifest)==46 and len(recipes)==46
    assert sum(r['kind']=='slides' for r in manifest)==22
    assert sum(r['kind']=='script' for r in manifest)==22
    assert sum(r['kind']=='reserve' for r in manifest)==2
    try:
        wrong=PBKDF2HMAC(algorithm=hashes.SHA256(),length=32,salt=base64.b64decode(meta['salt']),iterations=meta['iterations']).derive(material+b'wrong')
        v=(DIST/'protected/manifest.enc').read_bytes();AESGCM(wrong).decrypt(v[4:16],v[16:],None)
        raise AssertionError('Wrong password was accepted')
    except InvalidTag:pass
    result={'agendas':2,'participantHandouts':19,'packs':2,'decryptedSourceExact':46,'wrongPasswordRejected':True,'encryptedParts':len(cache),'verifiedAt':__import__('datetime').datetime.now(__import__('datetime').timezone.utc).isoformat()}
    (ROOT/'verification'/'materials-result.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result))
if __name__=='__main__':main()
